import logging
from django.contrib import admin
from django.urls import path, reverse
from django.shortcuts import render
from django.http import HttpResponseRedirect, HttpResponse
from django.contrib import messages
from django.utils.html import format_html
from django.db.models import Q
import openpyxl

from .models import (
    Course, Session, StudentEnrollment, CourseEnrollment,
    AcademyInstallment, SMSLog, Attendance, SessionMaterial, StudentIDCard,
    Exam, Question, Choice, MatchingPair, FillBlankAnswer, OrderingItem,
    ExamAttempt, StudentAnswer, ExpenseCategory, PaymentGateway, AccountingTransaction,
    DiscountCode, Teacher, StudentAccount, OrganizationMembership,
)
from apps.academy.admin_mixins import SchoolAdminMixin
from apps.notifications.sms import send_sms

logger = logging.getLogger(__name__)


# --- ۱. ارسال پیامک مستقیم با IPPanel ---

def send_profile_link_sms(receptor, student_name, course_title, profile_link):
    message = (
        f"{student_name} عزیز،\n"
        f"ثبت‌نام شما در دوره «{course_title}» تأیید شد.\n"
        f"مشاهده پروفایل و مدارک:\n{profile_link}"
    )
    return send_plain_sms(receptor, message)


def send_plain_sms(receptor, message_text):
    is_sent_status = send_sms(receptor, message_text)
    try:
        SMSLog.objects.create(
            receptor=receptor,
            message=message_text,
            sms_type='bulk_plain',
            is_sent=is_sent_status,
        )
    except Exception:
        logger.exception("Failed to write the SMS delivery log.")
    return is_sent_status


# --- ۲. کلاس‌های خطی (Inline) ---

class CourseEnrollmentInline(admin.TabularInline):
    model = CourseEnrollment
    extra = 1
    verbose_name = "ثبت‌نام دوره و وضعیت مالی"
    verbose_name_plural = "دوره‌های ثبت‌نام شده و وضعیت مالی"

class ChoiceInline(admin.TabularInline):
    model = Choice
    extra = 2
    verbose_name = "گزینه"
    verbose_name_plural = "گزینه‌ها"
    fields = ('text', 'image', 'is_correct', 'order')

class MatchingPairInline(admin.TabularInline):
    model = MatchingPair
    extra = 1
    verbose_name = "جفت"
    verbose_name_plural = "جفت‌های جورکردنی"

class FillBlankAnswerInline(admin.TabularInline):
    model = FillBlankAnswer
    extra = 1
    verbose_name = "پاسخ"
    verbose_name_plural = "پاسخ‌های جای خالی"

class OrderingItemInline(admin.TabularInline):
    model = OrderingItem
    extra = 2
    verbose_name = "آیتم"
    verbose_name_plural = "آیتم‌های ترتیب‌دهی"

class QuestionInline(admin.TabularInline):
    model = Question
    extra = 1
    verbose_name = "سوال"
    verbose_name_plural = "سوالات"
    fields = ('question_type', 'title', 'points', 'order', 'image', 'is_required')
    show_change_link = True


# --- ۳. مدیریت دوره‌های آموزشی ---

@admin.register(Course)
class CourseAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('title', 'code', 'start_date')
    search_fields = ('title', 'code')


# --- ۴. مدیریت جلسات دوره ---

@admin.register(Session)
class SessionAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'course__school'
    list_display = ('title', 'course', 'session_number', 'date', 'meeting_type', 'meeting_link_display')
    list_filter = ('course', 'meeting_type')
    search_fields = ('title', 'course__title')
    fieldsets = (
        ('اطلاعات جلسه', {
            'fields': ('course', 'title', 'session_number', 'date', 'start_time', 'end_time', 'description')
        }),
        ('کلاس آنلاین', {
            'fields': ('meeting_type', 'meeting_link', 'meeting_id', 'meeting_password'),
            'description': 'نوع کلاس آنلاین را انتخاب کنید. برای جیتسی لینک به صورت خودکار تولید می‌شود.',
        }),
    )

    def meeting_link_display(self, obj):
        link = obj.get_meeting_link()
        if link:
            return format_html('<a href="{}" target="_blank" style="color:#2563eb;font-weight:600;">ورود به کلاس</a>', link)
        return '—'
    meeting_link_display.short_description = 'لینک کلاس'


# --- ۵. پذیرش و ثبت‌نام هنرجویان اصلی ---

@admin.register(StudentEnrollment)
class StudentEnrollmentAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('first_name', 'last_name', 'phone_number', 'national_code', 'short_slug', 'is_documents_approved')
    list_filter = ('is_documents_approved',)
    search_fields = ('last_name', 'national_code', 'phone_number')
    inlines = [CourseEnrollmentInline]
    actions = ["send_bulk_sms_action"]

    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path('download-sample-excel/', self.admin_site.admin_view(self.download_sample_excel), name="download_sample_excel"),
            path('send-custom-sms-page/', self.admin_site.admin_view(self.send_custom_sms_page), name="send_custom_sms_page"),
            path('import-excel/', self.admin_site.admin_view(self.import_excel_action), name="import_excel_action"),
        ]
        return custom_urls + urls

    def download_sample_excel(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Suntech_Template"
        ws.views.sheetView[0].showGridLines = True
        headers = ["نام", "نام خانوادگی", "کد ملی", "شماره موبایل", "کد دوره", "مبلغ کل شهریه (تومان)", "مبلغ پرداخت شده (تومان)"]
        ws.append(headers)
        ws.append(["علی", "احمدی", "1271234567", "09123456789", "AI_BASE_01", 15000000, 5000000])
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=suntech_students_template.xlsx"
        wb.save(response)
        return response

    def import_excel_action(self, request):
        if request.method == "POST":
            excel_file = request.FILES.get("excel_file")
            if not excel_file or not excel_file.name.endswith('.xlsx'):
                messages.error(request, 'خطا: فرمت فایل باید حتما xlsx باشد.')
                return HttpResponseRedirect(".")
            wb = openpyxl.load_workbook(excel_file)
            worksheet = wb.active
            success_count = 0
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                first_name, last_name, national_code, phone_number, course_code, total, paid = row
                try:
                    course = Course.objects.get(code=course_code)
                    student, created = StudentEnrollment.objects.get_or_create(
                        national_code=str(national_code),
                        defaults={'first_name': first_name, 'last_name': last_name, 'phone_number': str(phone_number)}
                    )
                    if not CourseEnrollment.objects.filter(student=student, course=course).exists():
                        CourseEnrollment.objects.create(
                            student=student, course=course,
                            total_amount=int(total or 0), paid_amount=int(paid or 0),
                            payment_method='cash' if int(total or 0) == int(paid or 0) else 'installment'
                        )
                        success_count += 1
                        try:
                            relative_profile_url = reverse('academy:student_profile', kwargs={'slug': student.short_slug})
                        except Exception:
                            relative_profile_url = reverse('student_profile', kwargs={'slug': student.short_slug})
                        full_profile_link = request.build_absolute_uri(relative_profile_url)
                        full_name = f"{first_name} {last_name}"
                        send_profile_link_sms(str(phone_number), full_name, course.title, full_profile_link)
                except Course.DoesNotExist:
                    continue
            messages.success(request, f'تعداد {success_count} ثبت‌نام جدید با موفقیت اعمال و پیامک لینک پروفایل صادر شد.')
            return HttpResponseRedirect("..")
        context = self.admin_site.each_context(request)
        context.update({"title": "ورود از اکسل", "opts": self.model._meta})
        return render(request, "admin/academy_import_excel.html", context)

    @admin.display(description="ارسال پیامک دسته‌جمعی به هنرجویان انتخاب شده")
    def send_bulk_sms_action(self, request, queryset):
        selected_phones = [student.phone_number for student in queryset if student.phone_number]
        if not selected_phones:
            messages.warning(request, "هیچ هنرجویی با شماره موبایل معتبر انتخاب نشده است.")
            return HttpResponseRedirect(request.get_full_path())
        request.session['sms_receptors'] = selected_phones
        return HttpResponseRedirect("send-custom-sms-page/")

    def send_custom_sms_page(self, request):
        receptors = request.session.get('sms_receptors', [])
        if not receptors:
            messages.error(request, "شماره موبایلی یافت نشد. ابتدا هنرجویان را از لیست انتخاب کنید.")
            return HttpResponseRedirect("..")
        if request.method == "POST":
            message_text = request.POST.get("message_text")
            if message_text:
                success_count = 0
                for phone in receptors:
                    if send_plain_sms(phone, message_text):
                        success_count += 1
                if 'sms_receptors' in request.session:
                    del request.session['sms_receptors']
                messages.success(request, f"پیامک شما با موفقیت به {success_count} نفر از هنرجویان ارسال شد.")
                return HttpResponseRedirect("..")
        context = self.admin_site.each_context(request)
        context.update({
            "title": "ارسال پیامک دسته‌جمعی",
            "opts": self.model._meta,
            "receptor_count": len(receptors),
        })
        return render(request, "admin/academy_bulk_sms.html", context)

    def save_model(self, request, obj, form, change):
        is_new = obj.pk is None
        super().save_model(request, obj, form, change)
        if is_new and obj.phone_number:
            enrollment = CourseEnrollment.objects.filter(student=obj).last()
            course_title = enrollment.course.title if enrollment else "عمومی آکادمی"
            try:
                relative_profile_url = reverse('academy:student_profile', kwargs={'slug': obj.short_slug})
            except Exception:
                logger.exception("Could not reverse student profile URL for %s", obj.pk)
                relative_profile_url = f"/dashboard/students/{obj.short_slug}/"
            full_profile_link = request.build_absolute_uri(relative_profile_url)
            full_name = f"{obj.first_name} {obj.last_name}"
            send_profile_link_sms(obj.phone_number, full_name, course_title, full_profile_link)


# --- ۶. ثبت‌نام دوره (مستقل) ---

@admin.register(CourseEnrollment)
class CourseEnrollmentAdmin(admin.ModelAdmin):
    list_display = ('student', 'course', 'payment_method', 'total_amount', 'paid_amount')
    list_filter = ('payment_method', 'course')
    search_fields = ('student__last_name', 'student__phone_number', 'course__title')



# --- ۷. اقساط شهریه ---

@admin.register(AcademyInstallment)
class AcademyInstallmentAdmin(admin.ModelAdmin):
    list_display = ('enrollment', 'amount', 'due_date', 'status', 'paid_at')
    list_filter = ('status',)
    search_fields = ('enrollment__student__last_name', 'enrollment__course__title')



# --- ۸. لاگ پیامک‌ها ---

@admin.register(SMSLog)
class SMSLogAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('receptor', 'sms_type', 'is_sent', 'created_at')
    list_filter = ('sms_type', 'is_sent', 'created_at')
    search_fields = ('receptor', 'message')
    readonly_fields = ('receptor', 'message', 'sms_type', 'is_sent', 'created_at')


# --- ۹. حضور و غیاب ---

@admin.register(Attendance)
class AttendanceAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'session__course__school'
    list_display = ('student', 'session', 'status', 'recorded_at')
    list_filter = ('status', 'session__course')
    search_fields = ('student__last_name', 'student__phone_number', 'session__title')


# --- ۱۰. محتوای آموزشی ---

@admin.register(SessionMaterial)
class SessionMaterialAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'course__school'
    list_display = ('title', 'course', 'session', 'material_type', 'is_active')
    list_filter = ('material_type', 'is_active', 'course')
    search_fields = ('title', 'course__title', 'session__title')


# --- ۱۱. کارت شناسایی هنرجو ---

@admin.register(StudentIDCard)
class StudentIDCardAdmin(admin.ModelAdmin):
    list_display = ('student', 'card_number', 'is_valid', 'issued_at')
    list_filter = ('is_valid',)
    search_fields = ('student__last_name', 'card_number')



# --- ۱۲. آزمون‌ها ---

@admin.register(Exam)
class ExamAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'course__school'
    list_display = ('title', 'course', 'exam_type', 'total_questions', 'duration_minutes', 'is_active', 'status_display')
    list_filter = ('exam_type', 'is_active', 'course')
    search_fields = ('title', 'course__title')
    inlines = [QuestionInline]
    fieldsets = (
        ("اطلاعات اصلی", {
            'fields': ('course', 'title', 'description', 'exam_type')
        }),
        ("زمان‌بندی و نمره", {
            'fields': ('duration_minutes', 'pass_score_percent', 'attempts_allowed')
        }),
        ("تنظیمات پیشرفته", {
            'fields': ('shuffle_questions', 'shuffle_choices', 'show_result_immediately', 'show_answers_after', 'random_question_count', 'is_active', 'start_date', 'end_date')
        }),
    )


@admin.register(Question)
class QuestionAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'exam__course__school'
    list_display = ('title_preview', 'exam', 'question_type', 'points', 'order')
    list_filter = ('question_type', 'exam__course')
    search_fields = ('title', 'exam__title')
    fieldsets = (
        ("متن سوال", {
            'fields': ('exam', 'question_type', 'title', 'image', 'points', 'order', 'is_required')
        }),
        ("توضیحات", {
            'fields': ('explanation',),
            'classes': ('wide',),
        }),
    )

    def title_preview(self, obj):
        return obj.title[:60] + "..." if len(obj.title) > 60 else obj.title
    title_preview.short_description = "متن سوال"

    def get_inlines(self, request, obj):
        if obj and obj.question_type in ('multiple_choice', 'multiple_answer', 'true_false'):
            return [ChoiceInline]
        elif obj and obj.question_type == 'matching':
            return [MatchingPairInline]
        elif obj and obj.question_type == 'fill_blank':
            return [FillBlankAnswerInline]
        elif obj and obj.question_type == 'ordering':
            return [OrderingItemInline]
        return []


@admin.register(Choice)
class ChoiceAdmin(admin.ModelAdmin):
    list_display = ('text', 'question', 'is_correct', 'order')
    list_filter = ('is_correct', 'question__exam')
    search_fields = ('text', 'question__title')



@admin.register(MatchingPair)
class MatchingPairAdmin(admin.ModelAdmin):
    list_display = ('left_text', 'right_text', 'question', 'order')
    search_fields = ('left_text', 'right_text', 'question__title')



@admin.register(FillBlankAnswer)
class FillBlankAnswerAdmin(admin.ModelAdmin):
    list_display = ('answer_text', 'question', 'order')
    search_fields = ('answer_text', 'question__title')



@admin.register(OrderingItem)
class OrderingItemAdmin(admin.ModelAdmin):
    list_display = ('text', 'question', 'correct_order')
    search_fields = ('text', 'question__title')



@admin.register(ExamAttempt)
class ExamAttemptAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'exam__course__school'
    list_display = ('student', 'exam', 'attempt_number', 'status', 'score', 'max_score', 'is_passed', 'start_time')
    list_filter = ('status', 'is_passed', 'exam__course')
    search_fields = ('student__first_name', 'student__last_name', 'exam__title')
    readonly_fields = ('start_time', 'end_time', 'duration_display')

    def duration_display(self, obj):
        return obj.duration_display
    duration_display.short_description = "مدت زمان"


@admin.register(StudentAnswer)
class StudentAnswerAdmin(admin.ModelAdmin):
    list_display = ('attempt', 'question', 'is_correct', 'score', 'graded_at')
    list_filter = ('is_correct', 'graded_by')
    search_fields = ('attempt__student__last_name', 'question__title')



# --- ۱۳. دسته‌بندی هزینه‌ها ---

@admin.register(ExpenseCategory)
class ExpenseCategoryAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('name', 'description')
    search_fields = ('name',)


# --- ۱۴. درگاه‌های پرداخت ---

@admin.register(PaymentGateway)
class PaymentGatewayAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('name', 'gateway_type', 'card_number', 'account_holder', 'is_active')
    list_filter = ('gateway_type', 'is_active')
    search_fields = ('name', 'card_number')


# --- ۱۵. دفتر کل تراکنش‌های مالی ---

@admin.register(AccountingTransaction)
class AccountingTransactionAdmin(SchoolAdminMixin, admin.ModelAdmin):
    list_display = ('transaction_type', 'amount', 'payment_method', 'course', 'student', 'transaction_date', 'created_by')
    list_filter = ('transaction_type', 'payment_method', 'transaction_date')
    search_fields = ('description', 'receipt_number', 'student__last_name', 'course__title')
    readonly_fields = ('created_at', 'updated_at')



# --- ۱۶. تأیید مدارک ارسالی ---

class StudentDocument(StudentEnrollment):
    class Meta:
        proxy = True
        verbose_name = "تأیید مدرک هنرجو"
        verbose_name_plural = "تأیید مدارک ارسالی"

@admin.register(StudentDocument)
class StudentDocumentAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('first_name', 'last_name', 'national_code', 'is_documents_approved')
    list_filter = ('is_documents_approved',)
    search_fields = ('last_name', 'national_code')
    fields = ('first_name', 'last_name', 'avatar_3x4', 'national_card_img', 'identity_img', 'is_documents_approved')


# --- ۱۷. وضعیت پرداخت اقساط/شهریه ---

class TuitionPayment(CourseEnrollment):
    class Meta:
        proxy = True
        verbose_name = "وضعیت مالی شهریه"
        verbose_name_plural = "وضعیت پرداخت اقساط/شهریه"

@admin.register(TuitionPayment)
class TuitionPaymentAdmin(admin.ModelAdmin):
    list_display = ('get_student_first_name', 'get_student_last_name', 'course', 'total_amount', 'paid_amount', 'payment_method')
    list_filter = ('payment_method', 'course')
    search_fields = ('student__last_name', 'student__phone_number', 'course__title')

    def get_student_first_name(self, obj): return obj.student.first_name
    get_student_first_name.short_description = 'نام'

    def get_student_last_name(self, obj): return obj.student.last_name
    get_student_last_name.short_description = 'نام خانوادگی'



# --- ۱۸. صدور و چاپ کارت شناسایی ---

class IdCardGenerator(StudentEnrollment):
    class Meta:
        proxy = True
        verbose_name = "کارت شناسایی هنرجو"
        verbose_name_plural = "صدور و چاپ کارت شناسایی"

@admin.register(IdCardGenerator)
class IdCardGeneratorAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('first_name', 'last_name', 'get_courses', 'print_card_link')
    list_filter = ('is_documents_approved',)
    search_fields = ('last_name', 'national_code')

    def get_courses(self, obj):
        courses = CourseEnrollment.objects.filter(student_id=obj.id)
        if courses.exists():
            return ", ".join([ce.course.title for ce in courses])
        return "بدون دوره"
    get_courses.short_description = 'دوره‌های ثبت‌نامی'

    def print_card_link(self, obj):
        try:
            url = reverse('academy:student_id_card', kwargs={'slug': obj.short_slug})
        except Exception:
            url = f"/dashboard/students/{obj.short_slug}/id-card/"
        return format_html(
            '<a href="{}" target="_blank" class="inline-block bg-slate-800 hover:bg-slate-900 text-white font-bold py-1 px-3 rounded text-xs transition-colors">'
            'نمایش و چاپ کارت'
            '</a>',
            url
        )
    print_card_link.short_description = "عملیات صدور"


# --- ۱۹. کدهای تخفیف ---

@admin.register(DiscountCode)
class DiscountCodeAdmin(SchoolAdminMixin, admin.ModelAdmin):
    school_field = 'school'
    list_display = ('code', 'discount_type', 'discount_value', 'used_count', 'max_uses', 'is_active', 'created_at')
    list_filter = ('is_active', 'discount_type')
    search_fields = ('code', 'description')
    readonly_fields = ('used_count', 'created_at')


@admin.register(Teacher)
class TeacherAdmin(admin.ModelAdmin):
    list_display = ('user', 'phone_number', 'specialization', 'is_active', 'salary_per_session')
    list_filter = ('is_active',)
    search_fields = ('user__username', 'user__first_name', 'user__last_name', 'phone_number', 'national_code')


@admin.register(StudentAccount)
class StudentAccountAdmin(admin.ModelAdmin):
    list_display = ('user', 'enrollment')
    search_fields = ('user__username', 'enrollment__first_name', 'enrollment__last_name', 'enrollment__national_code')


@admin.register(OrganizationMembership)
class OrganizationMembershipAdmin(admin.ModelAdmin):
    list_display = ('user', 'organization')
    search_fields = ('user__username', 'organization__title')
    autocomplete_fields = ('user', 'organization')
