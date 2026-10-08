import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, get_object_or_404, redirect
from django.views import View
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Sum, Count, Q
from django.utils import timezone
from apps.academy.mixins import AdminRequiredMixin

from ..models import Course, StudentEnrollment, CourseEnrollment, AcademyInstallment
from ..models.accounting import AccountingTransaction
from ..models.attendance import Attendance


class ReportDashboardView(AdminRequiredMixin, View):
    def get(self, request):
        today = timezone.now().date()

        total_students = StudentEnrollment.objects.all().count()
        total_courses = Course.objects.all().count()
        total_enrollments = CourseEnrollment.objects.all().count()

        total_revenue = CourseEnrollment.objects.all().aggregate(t=Sum('paid_amount'))['t'] or 0
        total_outstanding = AcademyInstallment.objects.filter(
            status__in=['pending', 'overdue']
        ).aggregate(t=Sum('amount'))['t'] or 0

        students_by_course = CourseEnrollment.objects.all().values(
            'course__title', 'course__code'
        ).annotate(
            count=Count('id'),
            total_paid=Sum('paid_amount'),
            total_amount=Sum('total_amount'),
        ).order_by('-count')

        monthly_revenue = AccountingTransaction.objects.filter(
            transaction_type='income'
        ).extra(
            select={'year': "strftime('%%Y', transaction_date)",
                    'month': "strftime('%%m', transaction_date)"}
        ).values('year', 'month').annotate(
            total=Sum('amount'),
            count=Count('id')
        ).order_by('year', 'month')[:12]

        enrollment_stats = []
        for item in students_by_course:
            enrollment_stats.append({
                'title': item['course__title'],
                'code': item['course__code'],
                'count': item['count'],
                'total_paid': item['total_paid'] or 0,
                'total_amount': item['total_amount'] or 0,
            })

        recent_enrollments = CourseEnrollment.objects.select_related(
            'student', 'course'
        ).order_by('-enrolled_at')[:10]

        total_sessions = sum(c.sessions.count() for c in Course.objects.all())
        total_attendance = Attendance.objects.all().count()

        return render(request, 'academy/dashboard/reports/dashboard.html', {
            'total_students': total_students,
            'total_courses': total_courses,
            'total_enrollments': total_enrollments,
            'total_revenue': total_revenue,
            'total_outstanding': total_outstanding,
            'enrollment_stats': enrollment_stats,
            'monthly_revenue': list(monthly_revenue),
            'recent_enrollments': recent_enrollments,
            'total_sessions': total_sessions,
            'total_attendance': total_attendance,
        })


class CourseStudentsReportView(AdminRequiredMixin, View):
    def get(self, request):
        courses = Course.objects.all()
        course_id_raw = request.GET.get('course_id', '')
        try:
            course_id = int(course_id_raw) if str(course_id_raw).strip() else None
        except (TypeError, ValueError):
            course_id = None

        students = []
        selected_course = None
        if course_id is not None:
            selected_course = get_object_or_404(Course.objects.all(), id=course_id)
            enrollments = CourseEnrollment.objects.filter(
                course=selected_course
            ).select_related('student').order_by('student__first_name')

            for enroll in enrollments:
                student = enroll.student
                installments = AcademyInstallment.objects.filter(enrollment=enroll)
                total_installments = installments.count()
                paid_installments = installments.filter(status='paid').count()
                overdue_installments = installments.filter(status='overdue').count()

                attendance_count = Attendance.objects.filter(
                    session__course=selected_course,
                    student=student,
                    status='present'
                ).count()
                total_sessions = selected_course.sessions.count()

                students.append({
                    'student': student,
                    'enrollment': enroll,
                    'total_installments': total_installments,
                    'paid_installments': paid_installments,
                    'overdue_installments': overdue_installments,
                    'attendance_count': attendance_count,
                    'total_sessions': total_sessions,
                })

        return render(request, 'academy/dashboard/reports/course_students.html', {
            'courses': courses,
            'selected_course': selected_course,
            'students': students,
        })


class CourseStudentsExportView(AdminRequiredMixin, View):
    def get(self, request, course_id):
        course = get_object_or_404(Course.objects.all(), id=course_id)
        enrollments = CourseEnrollment.objects.filter(
            course=course
        ).select_related('student').order_by('student__first_name')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"هنرجویان {course.title}"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        headers = [
            "ردیف", "نام", "نام خانوادگی", "کد ملی", "شماره موبایل",
            "شهریه کل", "پرداختی", "بدهی", "نوع پرداخت",
            "تعداد اقساط", "اقساط پرداخت شده", "اقساط معوقه",
            "وضعیت",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for i, enroll in enumerate(enrollments, 1):
            student = enroll.student
            installments = AcademyInstallment.objects.filter(enrollment=enroll)
            total_inst = installments.count()
            paid_inst = installments.filter(status='paid').count()
            overdue_inst = installments.filter(status='overdue').count()

            status_text = "تسویه کامل" if enroll.is_fully_paid else (
                "در حال پرداخت" if enroll.paid_amount > 0 else "بدهکار"
            )

            row_data = [
                i,
                student.first_name,
                student.last_name,
                student.national_code,
                student.phone_number,
                enroll.total_amount,
                enroll.paid_amount,
                enroll.remaining_amount,
                enroll.get_payment_method_display(),
                total_inst,
                paid_inst,
                overdue_inst,
                status_text,
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=i + 1, column=col, value=value)
                cell.alignment = cell_alignment
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 25)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="students_{course.code}.xlsx"'
        wb.save(response)
        return response


class AllStudentsReportView(AdminRequiredMixin, View):
    def get(self, request):
        students = StudentEnrollment.objects.all().order_by('-created_at')

        course_raw = request.GET.get('course', '')
        search = request.GET.get('search', '')
        status = request.GET.get('status', '')
        if status not in ('paid', 'debtor', ''):
            status = ''

        try:
            course_id = int(course_raw) if str(course_raw).strip() else None
        except (TypeError, ValueError):
            course_id = None
        course = str(course_id) if course_id is not None else ''

        enrollments_qs = CourseEnrollment.objects.select_related('course')

        if course_id is not None:
            enrollments_qs = enrollments_qs.filter(course_id=course_id)

        student_data = []
        for s in students:
            if search:
                if search not in s.first_name and search not in s.last_name and search not in s.national_code and search not in s.phone_number:
                    continue

            s_enrollments = [e for e in enrollments_qs if e.student_id == s.id]
            if status == 'paid':
                s_enrollments = [e for e in s_enrollments if e.is_fully_paid]
            elif status == 'debtor':
                s_enrollments = [e for e in s_enrollments if not e.is_fully_paid]
            if not s_enrollments:
                continue

            total_enrolled = sum(e.total_amount for e in s_enrollments)
            total_paid = sum(e.paid_amount for e in s_enrollments)
            total_remaining = total_enrolled - total_paid

            course_names = ', '.join(e.course.title for e in s_enrollments)

            student_data.append({
                'student': s,
                'enrollments': s_enrollments,
                'course_names': course_names,
                'total_enrolled': total_enrolled,
                'total_paid': total_paid,
                'total_remaining': total_remaining,
                'enrollment_count': len(s_enrollments),
                'is_fully_paid': total_remaining == 0,
            })

        courses = Course.objects.all()

        # صفحه‌بندی
        from django.core.paginator import Paginator
        page_raw = request.GET.get('page', 1)
        try:
            page_number = int(page_raw)
        except (TypeError, ValueError):
            page_number = 1
        paginator = Paginator(student_data, 20)
        page_obj = paginator.get_page(page_number)

        return render(request, 'academy/dashboard/reports/all_students.html', {
            'student_data': page_obj,
            'page_obj': page_obj,
            'courses': courses,
            'filter_course': course,
            'filter_search': search,
            'filter_status': status,
        })


class AllStudentsExportView(AdminRequiredMixin, View):
    def get(self, request):
        students = StudentEnrollment.objects.all().order_by('-created_at')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "لیست هنرجویان"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        headers = [
            "ردیف", "نام", "نام خانوادگی", "کد ملی", "شماره موبایل",
            "تعداد دوره‌ها", "مجموع شهریه", "مجموع پرداختی", "بدهی کل", "وضعیت",
            "تاریخ ثبت",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        for i, s in enumerate(students, 1):
            enrollments = CourseEnrollment.objects.filter(student=s)
            total_enrolled = sum(e.total_amount for e in enrollments)
            total_paid = sum(e.paid_amount for e in enrollments)
            total_remaining = total_enrolled - total_paid
            status_text = "تسویه" if total_remaining == 0 else (
                "بدهکار" if total_remaining > 0 else "-"
            )

            row_data = [
                i,
                s.first_name,
                s.last_name,
                s.national_code,
                s.phone_number,
                enrollments.count(),
                total_enrolled,
                total_paid,
                total_remaining,
                status_text,
                s.created_at.strftime('%Y/%m/%d') if s.created_at else '',
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=i + 1, column=col, value=value)
                cell.alignment = cell_alignment
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 25)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="all_students.xlsx"'
        wb.save(response)
        return response
