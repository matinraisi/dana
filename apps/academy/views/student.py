import openpyxl
from django.shortcuts import render, get_object_or_404, redirect
from django.views import View
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, HttpResponseForbidden
from django.urls import reverse
from django.utils import timezone
from apps.academy.mixins import AdminRequiredMixin

from django.db.models import Q
from django.core.paginator import Paginator
from ..models import Course, StudentEnrollment, CourseEnrollment, AcademyInstallment
from ..models.accounting import AccountingTransaction, PaymentGateway
from ..services.sms_service import SmsService
from ..utils.date_helper import to_gregorian
from apps.users.models import User
from apps.academy.models import StudentAccount
from apps.academy.org import get_instance_organization

# محدودیت‌های آپلود فایل
ALLOWED_IMAGE_TYPES = ['image/jpeg', 'image/png', 'image/webp']
MAX_FILE_SIZE = 5 * 1024 * 1024  # ۵ مگابایت


def validate_image_file(file):
    """اعتبارسنجی فایل تصویر — نوع و اندازه"""
    if file.content_type not in ALLOWED_IMAGE_TYPES:
        return False, 'فایل تصویری مجاز نیست. فقط JPEG، PNG و WebP پذیرفته می‌شود.'
    if file.size > MAX_FILE_SIZE:
        return False, f'حجم فایل بیشتر از {MAX_FILE_SIZE // (1024*1024)} مگابایت است.'
    return True, None


def _ensure_student_user(student: StudentEnrollment, send_welcome_sms: bool = False, login_url: str = ''):
    """اگر هنرجو حساب کاربری ندارد، خودکار بسازد."""
    from apps.users.views import DEFAULT_PASSWORD
    already_exists = hasattr(student, 'user_account')
    if already_exists:
        return  # قبلاً ساخته شده
    username = f"s_{student.phone_number}"
    user, created = User.objects.get_or_create(
        username=username,
        defaults={
            'first_name': student.first_name,
            'last_name': student.last_name,
            'phone_number': student.phone_number,
            'role': 'STUDENT',
        }
    )
    if created or not user.has_usable_password():
        user.set_password(DEFAULT_PASSWORD)
    user.first_name = student.first_name
    user.last_name = student.last_name
    user.phone_number = student.phone_number
    user.role = 'STUDENT'
    user.save()
    StudentAccount.objects.get_or_create(user=user, defaults={'enrollment': student})
    # همیشه SMS بفرست اگر درخواست شده
    if send_welcome_sms and login_url:
        try:
            SmsService.notify_welcome(student, login_url)
        except Exception:
            pass


class StudentCreateView(AdminRequiredMixin, View):
    def get(self, request):
        courses = Course.objects.all()
        return render(request, 'academy/dashboard/student_create.html', {'courses': courses})

    def post(self, request):
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        national_code = request.POST.get('national_code')
        phone_number = request.POST.get('phone_number')
        course_id = request.POST.get('course')
        total_amount = request.POST.get('total_amount') or 0
        paid_amount = request.POST.get('paid_amount') or 0
        emergency_phone = request.POST.get('emergency_phone') or None
        field_of_study = request.POST.get('field_of_study') or None
        education_level = request.POST.get('education_level') or None
        birth_date = to_gregorian(request.POST.get('birth_date'))
        address = request.POST.get('address') or None

        if not all([first_name, last_name, national_code, phone_number, course_id]):
            messages.error(request, "خطا: پر کردن تمامی فیلدهای ستاره‌دار الزامی است.")
            return redirect('academy:student_create')

        try:
            course = Course.objects.get(id=course_id)

            duplicate_phone = StudentEnrollment.objects.filter(
                phone_number=str(phone_number)
            ).exclude(national_code=str(national_code)).first()

            if duplicate_phone:
                messages.error(request, f"این شماره تلفن قبلاً برای {duplicate_phone.full_name} ثبت شده است.")
                return redirect('academy:student_create')

            student, created = StudentEnrollment.objects.get_or_create(
                national_code=str(national_code),
                defaults={
                    'first_name': first_name, 'last_name': last_name,
                    'phone_number': str(phone_number), 'emergency_phone': emergency_phone,
                    'field_of_study': field_of_study, 'education_level': education_level,
                    'birth_date': birth_date or None, 'address': address,
                    'school': get_instance_organization(),
                }
            )

            if CourseEnrollment.objects.filter(student=student, course=course).exists():
                messages.warning(request, f"هنرجو {student} قبلاً در دوره {course.title} ثبت‌نام شده است.")
                return redirect('academy:dashboard_home')

            try:
                total = int(total_amount)
                paid = int(paid_amount)
            except (ValueError, TypeError):
                messages.error(request, "مبالغ وارد شده معتبر نیستند.")
                return redirect('academy:student_create')
            payment_method = 'cash' if total == paid and total > 0 else ('installment' if paid > 0 else 'unpaid')

            CourseEnrollment.objects.create(
                student=student, course=course,
                total_amount=total, paid_amount=paid,
                payment_method=payment_method
            )

            # ساخت خودکار حساب کاربری و ارسال پیامک خوش‌آمد
            login_url = request.build_absolute_uri(reverse('users:student_login'))
            _ensure_student_user(student, send_welcome_sms=True, login_url=login_url)

            # ارسال پیامک تأیید ثبت‌نام در دوره
            try:
                SmsService.notify_registration(student, course, request.build_absolute_uri(
                    reverse('academy:student_profile', kwargs={'slug': student.short_slug})
                ))
            except Exception:
                pass

            messages.success(request, f"ثبت‌نام هنرجو «{student.full_name}» با موفقیت انجام شد.")
            return redirect('academy:dashboard_home')

        except Course.DoesNotExist:
            messages.error(request, "دوره انتخاب شده یافت نشد.")
            return redirect('academy:student_create')
        except Exception as e:
            messages.error(request, f"خطایی رخ داد: {e}")
            return redirect('academy:student_create')


class StudentProfileView(LoginRequiredMixin, View):
    def get(self, request, slug):
        student = get_object_or_404(StudentEnrollment, short_slug=slug)
        # فقط خود دانش‌آموز، مدیر، یا سوپریوزر
        if not (request.user.is_superuser or
                request.user.is_admin_staff or
                (hasattr(request.user, 'student_account') and
                 request.user.student_account.enrollment_id == student.id)):
            return HttpResponseForbidden("دسترسی غیرمجاز.")
        return render(request, 'academy/student_profile.html', {'student': student})

    def post(self, request, slug):
        student = get_object_or_404(StudentEnrollment, short_slug=slug)
        # بررسی مالکیت
        if not (request.user.is_superuser or
                request.user.is_admin_staff or
                (hasattr(request.user, 'student_account') and
                 request.user.student_account.enrollment_id == student.id)):
            return HttpResponseForbidden("فقط مالک پروفایل یا مدیر می‌تواند مدارک را آپلود کند.")
        if request.FILES.get('avatar_3x4'):
            valid, err = validate_image_file(request.FILES['avatar_3x4'])
            if not valid:
                messages.error(request, err)
                return redirect('academy:student_profile', slug=slug)
            student.avatar_3x4 = request.FILES['avatar_3x4']
        if request.FILES.get('national_card_img'):
            valid, err = validate_image_file(request.FILES['national_card_img'])
            if not valid:
                messages.error(request, err)
                return redirect('academy:student_profile', slug=slug)
            student.national_card_img = request.FILES['national_card_img']
        if request.FILES.get('identity_img'):
            valid, err = validate_image_file(request.FILES['identity_img'])
            if not valid:
                messages.error(request, err)
                return redirect('academy:student_profile', slug=slug)
            student.identity_img = request.FILES['identity_img']
        student.save()
        messages.success(request, "مدارک با موفقیت بروزرسانی شد.")
        return redirect('academy:student_profile', slug=slug)


class StudentIdCardView(LoginRequiredMixin, View):
    def get(self, request, slug):
        student = get_object_or_404(StudentEnrollment, short_slug=slug)
        if not (request.user.is_superuser or
                request.user.is_admin_staff or
                (hasattr(request.user, 'student_account') and
                 request.user.student_account.enrollment_id == student.id)):
            return HttpResponseForbidden("دسترسی غیرمجاز.")
        return render(request, "academy/id_card_template.html", {"student": student})


class StudentEditView(AdminRequiredMixin, View):
    def get(self, request, slug):
        student = self.school_object_or_404(StudentEnrollment, short_slug=slug)
        courses = Course.objects.all()
        enrollments = CourseEnrollment.objects.filter(student=student)
        return render(request, 'academy/dashboard/student_edit.html', {
            'student': student, 'courses': courses, 'enrollments': enrollments,
        })

    def post(self, request, slug):
        student = self.school_object_or_404(StudentEnrollment, short_slug=slug)
        action = request.POST.get('action', 'save_info')

        # ─── حذف دوره ───
        if action == 'delete_enrollment':
            enroll_id = request.POST.get('delete_enrollment_id')
            enroll = get_object_or_404(CourseEnrollment, id=enroll_id, student=student)
            enroll.delete()
            messages.success(request, f"دوره «{enroll.course.title}» از هنرجو حذف شد.")
            return redirect('academy:student_edit', slug=student.short_slug)

        # ─── ویرایش مبالغ دوره ───
        if action == 'edit_enrollment':
            enroll_id = request.POST.get('edit_enrollment_id')
            enroll = get_object_or_404(CourseEnrollment, id=enroll_id, student=student)
            try:
                total = int(request.POST.get('edit_total_amount') or 0)
                paid = int(request.POST.get('edit_paid_amount') or 0)
            except (ValueError, TypeError):
                messages.error(request, "مبالغ وارد شده معتبر نیستند.")
                return redirect('academy:student_edit', slug=student.short_slug)
            enroll.total_amount = total
            enroll.paid_amount = paid
            enroll.payment_method = 'cash' if total == paid and total > 0 else ('installment' if paid > 0 else 'unpaid')
            enroll.save(update_fields=['total_amount', 'paid_amount', 'payment_method'])
            messages.success(request, f"مبالغ دوره «{enroll.course.title}» بروزرسانی شد.")
            return redirect('academy:student_edit', slug=student.short_slug)

        # ─── افزودن دوره جدید ───
        if action == 'add_course':
            course_id = request.POST.get('course')
            if not course_id:
                messages.error(request, "لطفاً یک دوره انتخاب کنید.")
                return redirect('academy:student_edit', slug=student.short_slug)
            try:
                course = Course.objects.get(id=course_id)
                if CourseEnrollment.objects.filter(student=student, course=course).exists():
                    messages.warning(request, f"هنرجو قبلاً در دوره «{course.title}» ثبت‌نام شده است.")
                    return redirect('academy:student_edit', slug=student.short_slug)
                try:
                    total = int(request.POST.get('total_amount') or 0)
                    paid = int(request.POST.get('paid_amount') or 0)
                except (ValueError, TypeError):
                    messages.error(request, "مبالغ وارد شده معتبر نیستند.")
                    return redirect('academy:student_edit', slug=student.short_slug)
                payment_method = 'cash' if total == paid and total > 0 else ('installment' if paid > 0 else 'unpaid')
                enroll = CourseEnrollment.objects.create(
                    student=student, course=course,
                    total_amount=total, paid_amount=paid,
                    payment_method=payment_method
                )
                if paid > 0:
                    default_gateway = PaymentGateway.objects.filter(
                        gateway_type__in=['pos', 'gateway', 'card', 'cash'], is_active=True
                    ).first()
                    AccountingTransaction.objects.create(
                        transaction_type='income',
                        amount=paid,
                        payment_method=payment_method,
                        payment_gateway=default_gateway,
                        course=course, student=student, enrollment=enroll,
                        description=f"پرداخت اولیه ثبت‌نام — {student}",
                        transaction_date=timezone.now().date(),
                        created_by=request.user,
                    )
                messages.success(request, f"دوره «{course.title}» با موفقیت به هنرجو اضافه شد.")
            except Course.DoesNotExist:
                messages.error(request, "دوره انتخاب شده یافت نشد.")
            return redirect('academy:student_edit', slug=student.short_slug)

        # ─── ذخیره اطلاعات فردی ───
        try:
            student.first_name = request.POST.get('first_name', student.first_name)
            student.last_name = request.POST.get('last_name', student.last_name)
            student.national_code = request.POST.get('national_code', student.national_code)
            student.phone_number = request.POST.get('phone_number', student.phone_number)
            student.emergency_phone = request.POST.get('emergency_phone') or None
            student.field_of_study = request.POST.get('field_of_study') or None
            student.education_level = request.POST.get('education_level') or None
            student.birth_date = to_gregorian(request.POST.get('birth_date'))
            student.address = request.POST.get('address') or None
            if not student.school_id:
                student.school = get_instance_organization()
            student.save()
            messages.success(request, f"اطلاعات هنرجو «{student.full_name}» با موفقیت بروزرسانی شد.")
        except Exception as e:
            messages.error(request, f"خطا در ذخیره اطلاعات: {e}")

        return redirect('academy:student_edit', slug=student.short_slug)


class StudentDeleteView(AdminRequiredMixin, View):
    def post(self, request, slug):
        student = self.school_object_or_404(StudentEnrollment, short_slug=slug)
        full_name = student.full_name
        try:
            if hasattr(student, 'user_account'):
                student.user_account.user.delete()
        except Exception:
            pass
        student.delete()
        messages.success(request, f"هنرجو «{full_name}» با موفقیت حذف شد.")
        return redirect('academy:dashboard_home')


class ImportExcelView(AdminRequiredMixin, View):
    def get(self, request):
        return render(request, 'academy/dashboard/import_excel.html')

    def post(self, request):
        excel_file = request.FILES.get("excel_file")
        if not excel_file or not excel_file.name.endswith('.xlsx'):
            messages.error(request, 'خطا: فرمت فایل باید xlsx باشد.')
            return redirect('academy:import_excel')

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            worksheet = wb.active
            success_count = 0
            courses = Course.objects.all()

            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not row or not row[2]:
                    continue
                first_name, last_name, national_code, phone_number, course_code = row[0], row[1], row[2], row[3], row[4]
                total = row[5] if len(row) > 5 else 0
                paid = row[6] if len(row) > 6 else 0
                field_of_study = row[7] if len(row) > 7 else None
                education_level = row[8] if len(row) > 8 else None
                emergency_phone = row[9] if len(row) > 9 else None
                birth_date = row[10] if len(row) > 10 else None
                address = row[11] if len(row) > 11 else None

                try:
                    course = courses.get(code=str(course_code).strip())
                    student, created = StudentEnrollment.objects.get_or_create(
                        national_code=str(national_code).strip(),
                        defaults={
                            'first_name': first_name, 'last_name': last_name,
                            'phone_number': str(phone_number).strip(),
                            'field_of_study': field_of_study, 'education_level': education_level,
                            'emergency_phone': str(emergency_phone).strip() if emergency_phone else None,
                            'birth_date': birth_date or None, 'address': address,
                            'school': get_instance_organization(),
                        }
                    )

                    if not CourseEnrollment.objects.filter(student=student, course=course).exists():
                        total_amt = int(total or 0)
                        paid_amt = int(paid or 0)
                        payment_method = 'cash' if total_amt == paid_amt and total_amt > 0 else ('installment' if paid_amt > 0 else 'unpaid')
                        CourseEnrollment.objects.create(
                            student=student, course=course,
                            total_amount=total_amt, paid_amount=paid_amt,
                            payment_method=payment_method
                        )
                        success_count += 1
                except (Course.DoesNotExist, Exception):
                    continue

            messages.success(request, f'تعداد {success_count} ثبت‌نام جدید با موفقیت اعمال شد.')
            return redirect('academy:dashboard_home')

        except Exception as e:
            messages.error(request, f'خطا در خواندن فایل اکسل: {e}')
            return redirect('academy:import_excel')


class StudentListView(AdminRequiredMixin, View):
    def get(self, request):
        courses = Course.objects.all()
        course_id = request.GET.get('course')
        search = request.GET.get('search', '').strip()

        students = StudentEnrollment.objects.prefetch_related('active_enrollments__course')

        if course_id:
            students = students.filter(active_enrollments__course_id=course_id)
        if search:
            students = students.filter(
                Q(first_name__icontains=search) | Q(last_name__icontains=search) |
                Q(phone_number__icontains=search) | Q(national_code__icontains=search)
            )
        students = students.order_by('-created_at')

        paginator = Paginator(students, 20)
        page = request.GET.get('page', 1)
        page_obj = paginator.get_page(page)

        context = {
            'courses': courses,
            'page_obj': page_obj,
            'selected_course': course_id,
            'search_query': search,
        }
        return render(request, 'academy/dashboard/student_list.html', context)


class DownloadSampleExcelView(AdminRequiredMixin, View):
    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ثبت‌نام هنرجویان"
        headers = ["نام", "نام خانوادگی", "کد ملی", "شماره موبایل", "کد دوره",
                   "شهریه کل (تومان)", "مبلغ پرداختی (تومان)", "رشته تحصیلی",
                   "مقطع تحصیلی", "تلفن اضطراری", "تاریخ تولد", "آدرس"]
        ws.append(headers)
        ws.append(["محمد", "رضایی", "0012345678", "09123456789", "PY-101",
                   "5000000", "2000000", "مهندسی کامپیوتر", "bachelor",
                   "09120001122", "1998-05-20", "تهران، خیابان آزادی"])
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = "attachment; filename=suntech_students_sample.xlsx"
        wb.save(response)
        return response
